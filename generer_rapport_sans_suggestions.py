import json
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows

def generer_rapport_sans_suggestions():
    # Lire les données
    with open('c:/Users/HeithemT/mabrouk-stock/exports/rapport_donnees_fraiches.json', 'r', encoding='utf-8') as f:
        data = json.load(f)
    
    wb = Workbook()
    
    # Supprimer la feuille par défaut
    wb.remove(wb.active)
    
    # Feuille 1: Articles Critiques
    ws_critiques = wb.create_sheet("Articles Critiques", 0)
    
    # En-têtes
    headers = ["Référence", "Code Article", "Saison", "Score", "Ventes/Semaine", "Stock Total", "Rupture Globale"]
    for col, header in enumerate(headers, 1):
        cell = ws_critiques.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")
    
    # Données
    row = 2
    for article in data.get('articlesCritiques', []):
        ws_critiques.cell(row=row, column=1, value=article.get('reference', ''))
        ws_critiques.cell(row=row, column=2, value=article.get('codeArticle', ''))
        ws_critiques.cell(row=row, column=3, value=article.get('saison', ''))
        ws_critiques.cell(row=row, column=4, value=article.get('score', 0))
        ws_critiques.cell(row=row, column=5, value=article.get('ventesParSemaine', 0))
        ws_critiques.cell(row=row, column=6, value=article.get('stockTotal', 0))
        ws_critiques.cell(row=row, column=7, value="OUI" if article.get('ruptureGlobale') else "NON")
        
        # Colorier les lignes selon le score
        score = article.get('score', 0)
        if score >= 10:
            fill_color = "FFFF0000"  # Rouge
        elif score >= 5:
            fill_color = "FFFFA500"  # Orange
        else:
            fill_color = "FFFFFF00"  # Jaune
            
        for col in range(1, 8):
            ws_critiques.cell(row=row, column=col).fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
        
        row += 1
    
    # Ajuster les colonnes
    for col in ws_critiques.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws_critiques.column_dimensions[column].width = adjusted_width
    
    # Feuille 2: Résumé
    ws_resume = wb.create_sheet("Résumé", 1)
    
    resume_data = [
        ["Période", data.get('periode', {}).get('label', '')],
        ["Date Génération", data.get('dateGeneration', '')],
        ["Articles Critiques", len(data.get('articlesCritiques', []))],
        ["Articles Faibles", len(data.get('articlesFaibles', []))],
        ["Total Suggestions", 0],
        ["", ""],
        ["⚠️ PROBLÈME DÉTECTÉ", ""],
        ["Aucun stock disponible", "pour les transferts"],
        ["Action requise", "Commande fournisseur"]
    ]
    
    for row_idx, (key, value) in enumerate(resume_data, 1):
        ws_resume.cell(row=row_idx, column=1, value=key)
        ws_resume.cell(row=row_idx, column=2, value=value)
        
        if "PROBLÈME" in str(key):
            ws_resume.cell(row=row_idx, column=1).font = Font(bold=True, color="FF0000")
            ws_resume.cell(row=row_idx, column=1).fill = PatternFill(start_color="FFFF0000", end_color="FFFF0000", fill_type="solid")
    
    # Sauvegarder
    output_file = "c:/Users/HeithemT/mabrouk-stock/exports/rapport_avec_donnees.xlsx"
    wb.save(output_file)
    print(f"Rapport généré: {output_file}")
    print(f"Articles critiques: {len(data.get('articlesCritiques', []))}")
    print(f"Aucune suggestion générée - stock insuffisant")

if __name__ == "__main__":
    generer_rapport_sans_suggestions()
