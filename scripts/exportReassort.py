import json, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime

# Lire depuis fichier argument ou stdin
if len(sys.argv) >= 3:
    with open(sys.argv[2], 'r', encoding='utf-8') as f:
        data = json.load(f)
    output_path = sys.argv[1]
else:
    data = json.loads(sys.stdin.read())
    output_path = sys.argv[1] if len(sys.argv) > 1 else "/tmp/reassort.xlsx"

today = datetime.now().strftime("%d/%m/%Y")

RED_BG="FFEBEE"; RED_FG="C62828"; ORANGE_BG="FFF8E1"; ORANGE_FG="E65100"
GREEN_BG="E8F5E9"; GREEN_FG="2E7D32"; BLUE_BG="E3F2FD"; BLUE_FG="1565C0"
HEADER_BG="1E293B"; HEADER_FG="F59E0B"; GREY_BG="F8FAFC"

def side(): return Side(style='thin', color='E2E8F0')
def brd(): return Border(left=side(), right=side(), top=side(), bottom=side())

def hdr(cell, bg=HEADER_BG, fg=HEADER_FG, sz=11):
    cell.font = Font(bold=True, color=fg, size=sz, name='Arial')
    cell.fill = PatternFill('solid', start_color=bg)
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cell.border = brd()

def sty(cell, bg="FFFFFF", fg="1E293B", bold=False, center=False):
    cell.font = Font(color=fg, size=10, name='Arial', bold=bold)
    cell.fill = PatternFill('solid', start_color=bg)
    cell.alignment = Alignment(horizontal='center' if center else 'left', vertical='center', wrap_text=True)
    cell.border = brd()

wb = Workbook()

# ── FEUILLE 1 : TRANSFERTS ────────────────────────────────
ws1 = wb.active
ws1.title = "Transferts à faire"
ws1.sheet_view.showGridLines = False
ws1.row_dimensions[1].height = 45
ws1.row_dimensions[2].height = 22
ws1.row_dimensions[3].height = 35

critique_count = len(data.get('critique', []))
faible_count = len(data.get('faible', []))
all_items = [(x,'CRITIQUE') for x in data.get('critique',[])] + [(x,'FAIBLE') for x in data.get('faible',[])]
total_t = sum(len(x.get('suggestions',[])) for x,_ in all_items)

ws1.merge_cells('A1:I1')
t=ws1['A1']; t.value=f"MABROUK — Plan de Transferts   |   Généré le {today}"
t.font=Font(bold=True,color=HEADER_FG,size=14,name='Arial')
t.fill=PatternFill('solid',start_color=HEADER_BG)
t.alignment=Alignment(horizontal='center',vertical='center')

ws1.merge_cells('A2:I2')
s=ws1['A2']; s.value=f"Critiques : {critique_count}   |   Faibles : {faible_count}   |   Total transferts : {total_t}"
s.font=Font(color="FFFFFF",size=10,name='Arial')
s.fill=PatternFill('solid',start_color="334155")
s.alignment=Alignment(horizontal='center',vertical='center')

hdrs=['Urgence','Code-barres','Code Article','Saison','Commander ?','De (Donneur)','Vers (Receveur)','Quantité','Raison']
for col,h in enumerate(hdrs,1):
    hdr(ws1.cell(row=3,column=col,value=h))

for i,w in enumerate([12,18,13,10,20,20,20,10,45],1):
    ws1.column_dimensions[get_column_letter(i)].width=w

SAISONS_ACT=['25H','25E','26E']
row=4
for item,urgence in all_items:
    saison=item.get('saison','')
    actuel=saison.strip().upper() in SAISONS_ACT if saison else False
    commander="OUI - Commander" if actuel else "NON - Ancienne saison"
    for s in item.get('suggestions',[]):
        ws1.row_dimensions[row].height=22
        urg_bg=RED_BG if urgence=='CRITIQUE' else ORANGE_BG
        urg_fg=RED_FG if urgence=='CRITIQUE' else ORANGE_FG
        bg_row="FFFFFF" if row%2==0 else GREY_BG
        vals=[urgence,item.get('reference',''),item.get('codeArticle',''),saison,commander,s.get('de',''),s.get('vers',''),s.get('quantite',0),s.get('raison','')]
        for col,val in enumerate(vals,1):
            c=ws1.cell(row=row,column=col,value=val)
            if col==1: sty(c,bg=urg_bg,fg=urg_fg,bold=True,center=True)
            elif col==5: sty(c,bg=GREEN_BG if actuel else ORANGE_BG,fg=GREEN_FG if actuel else ORANGE_FG,center=True)
            elif col==8: sty(c,bg=urg_bg,fg=urg_fg,bold=True,center=True)
            elif col==6 and s.get('de')=='SITE CENTRALE': sty(c,bg=BLUE_BG,fg=BLUE_FG,bold=True)
            else: sty(c,bg=bg_row)
        row+=1

ws1.merge_cells(f'A{row}:G{row}')
tc=ws1[f'A{row}']; tc.value="TOTAL UNITES A TRANSFERER"
tc.font=Font(bold=True,color=HEADER_FG,size=11,name='Arial')
tc.fill=PatternFill('solid',start_color=HEADER_BG)
tc.alignment=Alignment(horizontal='right',vertical='center')
ws1.row_dimensions[row].height=25
qc=ws1.cell(row=row,column=8,value=f'=SUM(H4:H{row-1})')
hdr(qc)
ws1.cell(row=row,column=9).fill=PatternFill('solid',start_color=HEADER_BG)
ws1.freeze_panes='A4'

# ── FEUILLE 2 : RÉSUMÉ PAR BOUTIQUE ──────────────────────
ws2=wb.create_sheet("Resume par boutique")
ws2.sheet_view.showGridLines=False
ws2.row_dimensions[1].height=40
ws2.row_dimensions[2].height=32

ws2.merge_cells('A1:E1')
t2=ws2['A1']; t2.value=f"Resume des transferts par boutique   |   {today}"
t2.font=Font(bold=True,color=HEADER_FG,size=13,name='Arial')
t2.fill=PatternFill('solid',start_color=HEADER_BG)
t2.alignment=Alignment(horizontal='center',vertical='center')

boutiques_recoit={}; boutiques_donne={}
for item,_ in all_items:
    for s in item.get('suggestions',[]):
        de=s.get('de',''); vers=s.get('vers',''); qte=s.get('quantite',0)
        boutiques_donne[de]=boutiques_donne.get(de,0)+qte
        boutiques_recoit[vers]=boutiques_recoit.get(vers,0)+qte

for col,h in enumerate(['Boutique','Unites a recevoir','Unites a donner','Balance','Role'],1):
    hdr(ws2.cell(row=2,column=col,value=h))
for i,w in enumerate([25,18,18,12,20],1):
    ws2.column_dimensions[get_column_letter(i)].width=w

toutes=sorted(set(list(boutiques_recoit.keys())+list(boutiques_donne.keys())))
r2=3
for b in toutes:
    recoit=boutiques_recoit.get(b,0); donne=boutiques_donne.get(b,0); balance=recoit-donne
    role="Receveur" if recoit>donne else ("Donneur" if donne>recoit else "Neutre")
    bg=RED_BG if recoit>donne else (GREEN_BG if donne>recoit else "FFFFFF")
    ws2.row_dimensions[r2].height=22
    for col,val in enumerate([b,recoit,donne,balance,role],1):
        c=ws2.cell(row=r2,column=col,value=val)
        sty(c,bg=bg if col>1 else "FFFFFF",center=(col>1))
        if col==1: c.font=Font(bold=True,size=10,name='Arial',color="1E293B"); c.border=brd()
    r2+=1

ws2.row_dimensions[r2].height=25
for col in range(1,6):
    c=ws2.cell(row=r2,column=col)
    if col==1: c.value="TOTAL"
    elif col==2: c.value=f'=SUM(B3:B{r2-1})'
    elif col==3: c.value=f'=SUM(C3:C{r2-1})'
    elif col==4: c.value=f'=SUM(D3:D{r2-1})'
    hdr(c)
ws2.freeze_panes='A3'

# ── FEUILLE 3 : ARTICLES OK ───────────────────────────────
ws3=wb.create_sheet("Articles OK")
ws3.sheet_view.showGridLines=False
ws3.row_dimensions[1].height=35
ws3.row_dimensions[2].height=30
ws3.column_dimensions['A'].width=20
ws3.column_dimensions['B'].width=15

ws3.merge_cells('A1:B1')
t3=ws3['A1']; t3.value=f"Articles bien distribues ({len(data.get('ok',[]))}) — {today}"
t3.font=Font(bold=True,color="FFFFFF",size=12,name='Arial')
t3.fill=PatternFill('solid',start_color="166534")
t3.alignment=Alignment(horizontal='center',vertical='center')

for col,h in enumerate(['Code-barres','Statut'],1):
    hdr(ws3.cell(row=2,column=col,value=h),bg="166534")

for i,ref in enumerate(data.get('ok',[]),3):
    ws3.row_dimensions[i].height=20
    c1=ws3.cell(row=i,column=1,value=ref)
    sty(c1,bg="F0FDF4" if i%2==0 else "FFFFFF")
    c1.font=Font(size=10,name='Courier New',color="1E293B"); c1.border=brd()
    c2=ws3.cell(row=i,column=2,value="OK")
    sty(c2,bg=GREEN_BG,fg=GREEN_FG,bold=True,center=True)

wb.save(output_path)