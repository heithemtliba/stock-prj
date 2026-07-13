import json, sys, requests
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime

# ── ARTICLES A EXCLURE (sacs d'emballage Mabrouk) ─────────────────────────
ARTICLES_EXCLUS = {'91272', '91273', '91274'}

# ── SAISONS ────────────────────────────────────────────────────────────────
SAISONS_ACT             = ['25H', '25E', '26E']
SAISONS_NOUVELLE_COLL   = {'26E', '25H'}

# ── SEUILS DE TRANSFERT PAR ZONE ──────────────────────────────────────────
STORES_GRAND_TUNIS = {'002','005','014','015','016','019','021','024','030','031','033'}
STORES_CENTRE      = {'029'}   # Sousse
STORES_NORD        = {'011'}   # Nabeul
STORES_SUD         = {'009','032'}  # Sfax, Sfax Mall

def seuil_transfert(id_donneur, id_receveur):
    id_d = str(id_donneur); id_r = str(id_receveur)
    if id_d in STORES_SUD or id_r in STORES_SUD:    return 8
    if id_d in STORES_CENTRE or id_r in STORES_CENTRE: return 5
    if id_d in STORES_NORD or id_r in STORES_NORD:   return 5
    return 1
# ── STOCK TEMPS REEL VIA API NODE ─────────────────────────────────────────
API_BASE = "http://localhost:3002"

def get_stock_article(code_article):
    return {}  # Stock pre-calcule cote Node
    try:
        r = requests.get(f"{API_BASE}/stock-article/{code_article}", timeout=15)
        if r.status_code == 200:
            d = r.json()
            if d.get('success'):
                return {str(s['storeId']): int(float(s.get('stock', 0)))
                        for s in d.get('stockParBoutique', [])}
    except Exception:
        pass
    return {}

# ── NOMS BOUTIQUES ─────────────────────────────────────────────────────────
STORE_NAMES = {
    '002':'MENZAH','005':'JAMEL ABDENNACEUR','009':'SFAX A',
    '011':'NABEUL','014':'CARREFOUR','015':'ZEPHYR',
    '016':'LAFAYETTE','019':'E-Boutique','021':'TUNISIA MALL',
    '024':'GEANT 3','029':'SOUSSE SLIM CENTER','030':'AZUR CITY',
    '031':'SOUKRA','032':'MALL OF SFAX','033':'LAC Premium'
}

# ── LECTURE DES DONNEES ────────────────────────────────────────────────────
if len(sys.argv) >= 3:
    with open(sys.argv[2], 'r', encoding='utf-8') as f:
        data = json.load(f)
    output_path = sys.argv[1]
else:
    data = json.loads(sys.stdin.read())
    output_path = sys.argv[1] if len(sys.argv) > 1 else "/tmp/rapport.xlsx"

today   = datetime.now().strftime("%d/%m/%Y")
now_str = datetime.now().strftime("%d/%m/%Y %H:%M")

# ── FILTRAGE GLOBAL DES SACS ───────────────────────────────────────────────
def filtrer_articles(liste, cle='codeArticle'):
    return [a for a in liste if str(a.get(cle,'')).strip() not in ARTICLES_EXCLUS]

periode     = data.get('periode', {})
resume      = data.get('resumeUrgences', {})
critiques   = filtrer_articles(data.get('articlesCritiques', []))
faibles     = filtrer_articles(data.get('articlesFaibles', []))
ventesHebdo = data.get('ventesHebdo', {})
topArticles = filtrer_articles(data.get('topArticles', []), cle='code_article')
inactifs    = filtrer_articles(data.get('articlesInactifs', []), cle='code_article')
stock_par_article = data.get('stockParArticle', {})  # Stock temps reel pre-calcule
prix_articles    = data.get('prixArticles', {})  # Prix TTC par code article
rupturesNC  = filtrer_articles(data.get('rupturesNouvelleCollection', []), cle='code_article')
scores_mag  = data.get('scoresMagasins', [])
noms_art    = data.get('nomsArticles', {})  # { "22064": "Nom article", ... }

# ── STYLES ─────────────────────────────────────────────────────────────────
RED_BG   = "FFEBEE"; RED_FG   = "C62828"
ORANGE_BG= "FFF8E1"; ORANGE_FG= "E65100"
GREEN_BG = "E8F5E9"; GREEN_FG = "2E7D32"
BLUE_BG  = "E3F2FD"; BLUE_FG  = "1565C0"
PURPLE_BG= "F3E5F5"; PURPLE_FG= "6A1B9A"
GREY_BG  = "F8FAFC"
HEADER_BG= "1E293B"; HEADER_FG= "F59E0B"
SUB_BG   = "334155"; SUB_FG   = "FFFFFF"

def side():  return Side(style='thin', color='E2E8F0')
def brd():   return Border(left=side(), right=side(), top=side(), bottom=side())

def hdr(cell, bg=HEADER_BG, fg=HEADER_FG, sz=11, left=False):
    cell.font      = Font(bold=True, color=fg, size=sz, name='Arial')
    cell.fill      = PatternFill('solid', start_color=bg)
    cell.alignment = Alignment(horizontal='left' if left else 'center', vertical='center', wrap_text=True)
    cell.border    = brd()

def sty(cell, bg="FFFFFF", fg="1E293B", bold=False, center=False, sz=10):
    cell.font      = Font(color=fg, size=sz, name='Arial', bold=bold)
    cell.fill      = PatternFill('solid', start_color=bg)
    cell.alignment = Alignment(horizontal='center' if center else 'left', vertical='center', wrap_text=True)
    cell.border    = brd()

def titre_principal(ws, texte, nb_cols, row=1, h=45):
    ws.merge_cells(f'A{row}:{get_column_letter(nb_cols)}{row}')
    c = ws.cell(row=row, column=1, value=texte)
    c.font      = Font(bold=True, color=HEADER_FG, size=14, name='Arial')
    c.fill      = PatternFill('solid', start_color=HEADER_BG)
    c.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[row].height = h

def sous_titre(ws, texte, nb_cols, row=2, h=22):
    ws.merge_cells(f'A{row}:{get_column_letter(nb_cols)}{row}')
    c = ws.cell(row=row, column=1, value=texte)
    c.font      = Font(color=SUB_FG, size=10, name='Arial')
    c.fill      = PatternFill('solid', start_color=SUB_BG)
    c.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[row].height = h

wb = Workbook()

# ══════════════════════════════════════════════════════════════════════════
# FEUILLE 1 — TABLEAU DE BORD
# ══════════════════════════════════════════════════════════════════════════
ws1 = wb.active
ws1.title = "Tableau de bord"
ws1.sheet_view.showGridLines = False
titre_principal(ws1, f"MABROUK DIFFUSION — Rapport Hebdomadaire   |   {now_str}   |   Periode : {periode.get('label','')}", 6)
sous_titre(ws1, f"Analyse sur {periode.get('jours',180)} jours   |   Semaine du {today}", 6)

ws1.merge_cells('A3:F3')
c = ws1.cell(row=3, column=1, value="INDICATEURS CLES DE LA SEMAINE")
c.font = Font(bold=True, color=HEADER_FG, size=11, name='Arial')
c.fill = PatternFill('solid', start_color=SUB_BG)
c.alignment = Alignment(horizontal='center', vertical='center')
ws1.row_dimensions[3].height = 28

kpis = [
    ("Articles CRITIQUE", resume.get('nbCritique',0), RED_BG, RED_FG),
    ("Articles FAIBLE",   resume.get('nbFaible',0),   ORANGE_BG, ORANGE_FG),
    ("Articles OK",       resume.get('nbOk',0),        GREEN_BG, GREEN_FG),
    ("Transferts a faire",resume.get('totalTransferts',0), BLUE_BG, BLUE_FG),
    ("Ventes cette sem.", ventesHebdo.get('cetteSemaine',0), GREY_BG, "1E293B"),
    ("Evolution vs S-1",  f"{ventesHebdo.get('evolution',0):+}%", PURPLE_BG, PURPLE_FG),
]
for col, (label, valeur, bg, fg) in enumerate(kpis, 1):
    ws1.column_dimensions[get_column_letter(col)].width = 20
    for r, v, sz in [(4, label, 9), (5, valeur, 22)]:
        ws1.row_dimensions[r].height = 22 if r == 4 else 40
        c = ws1.cell(row=r, column=col, value=v)
        c.font = Font(bold=True, color=fg, size=sz, name='Arial')
        c.fill = PatternFill('solid', start_color=bg)
        c.alignment = Alignment(horizontal='center', vertical='center')
        c.border = brd()
    cd = ws1.cell(row=6, column=col)
    cd.fill = PatternFill('solid', start_color=bg); cd.border = brd()
    ws1.row_dimensions[6].height = 22

# TOP 5 critiques
ws1.merge_cells('A8:F8')
c8 = ws1.cell(row=8, column=1, value="TOP 5 ARTICLES CRITIQUE — ACTION IMMEDIATE REQUISE")
c8.font = Font(bold=True, color=RED_FG, size=11, name='Arial')
c8.fill = PatternFill('solid', start_color=RED_BG)
c8.alignment = Alignment(horizontal='center', vertical='center')
ws1.row_dimensions[8].height = 28

for col, h in enumerate(['Code Article','Nom Article','Saison','Score','Ventes/sem','Transferts suggeres'], 1):
    hdr(ws1.cell(row=9, column=col, value=h))
ws1.row_dimensions[9].height = 22

for i, art in enumerate(critiques[:5]):
    row = 10 + i
    ws1.row_dimensions[row].height = 22
    code = str(art.get('codeArticle',''))
    bg   = RED_BG if i % 2 == 0 else "FFCDD2"
    for col, val in enumerate([
        code, noms_art.get(code,''), art.get('saison',''),
        art.get('score',0), art.get('ventesParSemaine',0), len(art.get('suggestions',[]))
    ], 1):
        c = ws1.cell(row=row, column=col, value=val)
        if col == 4: sty(c, bg=bg, fg=RED_FG, bold=True, center=True)
        else: sty(c, bg=bg, fg=RED_FG)

# TOP 10 vendus
ws1.merge_cells('A16:F16')
c16 = ws1.cell(row=16, column=1, value="TOP 10 ARTICLES LES PLUS VENDUS — CETTE SEMAINE")
c16.font = Font(bold=True, color=GREEN_FG, size=11, name='Arial')
c16.fill = PatternFill('solid', start_color=GREEN_BG)
c16.alignment = Alignment(horizontal='center', vertical='center')
ws1.row_dimensions[16].height = 28

for col, h in enumerate(['Rang','Code Article','Nom Article','Saison','Unites vendues','Statut saison'], 1):
    hdr(ws1.cell(row=17, column=col, value=h))

for i, art in enumerate(topArticles, 1):
    row = 17 + i
    ws1.row_dimensions[row].height = 22
    code   = str(art.get('code_article',''))
    saison = art.get('saison','')
    actuel = saison.strip().upper() in SAISONS_ACT
    bg     = GREEN_BG if i % 2 == 0 else "FFFFFF"
    sty(ws1.cell(row=row, column=1, value=i), bg=bg, center=True, bold=(i<=3))
    sty(ws1.cell(row=row, column=2, value=code), bg=bg, bold=(i<=3))
    sty(ws1.cell(row=row, column=3, value=noms_art.get(code,'')), bg=bg)
    sty(ws1.cell(row=row, column=4, value=saison), bg=bg, center=True)
    sty(ws1.cell(row=row, column=5, value=art.get('total_vendu',0)), bg=bg, center=True, bold=True)
    sc = ws1.cell(row=row, column=6, value="Saison actuelle" if actuel else "Ancienne saison")
    sty(sc, bg=GREEN_BG if actuel else ORANGE_BG, fg=GREEN_FG if actuel else ORANGE_FG, center=True)

ws1.freeze_panes = 'A4'

# ══════════════════════════════════════════════════════════════════════════
# FEUILLE 2 — ARTICLES CRITIQUE + TRANSFERTS URGENTS (amelioree)
# Colonnes: Code | Nom | Saison | Score Donneur | Donneur | Stock Donneur
#           | Score Receveur | Receveur | Stock Receveur | Qte | Jours | Seuil | Consigne
# ══════════════════════════════════════════════════════════════════════════
ws2 = wb.create_sheet("Critique - Transferts urgents")
ws2.sheet_view.showGridLines = False
nb_cols2 = 13

titre_principal(ws2, f"ARTICLES CRITIQUE — TRANSFERTS URGENTS   |   {today}", nb_cols2)
sous_titre(ws2, f"{len(critiques)} article(s) critique(s)   |   Action requise sous 24-48h   |   Seuils transport appliques", nb_cols2)

hdrs2 = ['Code Article','Nom Article','Saison',
         'Score Donneur','De (Donneur)','Stock Donneur',
         'Score Receveur','Vers (Receveur)','Stock Receveur',
         'Quantite','Jours stock receveur','Seuil min.','Consigne']
for col, h in enumerate(hdrs2, 1):
    hdr(ws2.cell(row=3, column=col, value=h))
ws2.row_dimensions[3].height = 28

for i, w in enumerate([14,28,10,13,24,13,13,24,13,10,18,10,40], 1):
    ws2.column_dimensions[get_column_letter(i)].width = w

row = 4
for art in critiques:
    code   = str(art.get('codeArticle',''))
    nom    = noms_art.get(code,'')
    saison = art.get('saison','')
    actuel = saison.strip().upper() in SAISONS_ACT
    stock_reel = stock_par_article.get(str(code), {})
    store_info = {b.get('storeName',''): b for b in art.get('analyse',[])}

    for s in art.get('suggestions', []):
        nd = s.get('de','');  nr = s.get('vers','')
        bd = store_info.get(nd, {}); br = store_info.get(nr, {})
        # Utiliser les IDs déjà présents dans les suggestions (fiable même si 'analyse' est filtré)
        id_d = str(s.get('deId') or bd.get('storeId',''))
        id_r = str(s.get('versId') or br.get('storeId',''))
        # Le backend ne fournit pas de "score" par boutique dans 'analyse' -> fallback sur score article
        sc_article = art.get('score', 0)
        sc_d = s.get('scoreDonneur', sc_article)
        sc_r = s.get('scoreReceveur', sc_article)
        jours_r = br.get('joursStock', 0)
        stk_d = max(0, stock_reel.get(id_d, 0) if isinstance(stock_reel.get(id_d, 0), (int,float)) else 0)
        stk_r = max(0, stock_reel.get(id_r, 0) if isinstance(stock_reel.get(id_r, 0), (int,float)) else 0)
        seuil = seuil_transfert(id_d, id_r)
        quantite = s.get('quantite', 0)
        seuil = seuil_transfert(id_d, id_r)
        if isinstance(quantite, (int, float)) and quantite < seuil:
            continue
        bg = RED_BG if row % 2 == 0 else "FFCDD2"

        vals = [code, nom, saison, sc_d, nd, stk_d, sc_r, nr, stk_r,
                quantite, jours_r, seuil,
                "Commander si saison actuelle" if actuel else "Redistribuer uniquement"]

        for col, val in enumerate(vals, 1):
            c = ws2.cell(row=row, column=col, value=val)
            if   col == 1:  sty(c, bg=bg, fg=RED_FG, bold=True)
            elif col == 2:  sty(c, bg=bg, fg="1E293B")
            elif col == 4:  sty(c, bg=BLUE_BG, fg=BLUE_FG, bold=True, center=True)   # score donneur
            elif col == 5:  sty(c, bg=BLUE_BG if nd=='SITE CENTRALE' else bg,
                                fg=BLUE_FG if nd=='SITE CENTRALE' else RED_FG,
                                bold=(nd=='SITE CENTRALE'))
            elif col == 6:  sty(c, bg=BLUE_BG, fg=BLUE_FG, bold=True, center=True)   # stock donneur
            elif col == 7:  sty(c, bg=GREEN_BG, fg=GREEN_FG, bold=True, center=True)  # score receveur
            elif col == 8:  sty(c, bg=GREEN_BG, fg=GREEN_FG)
            elif col == 9:  sty(c, bg=GREEN_BG, fg=GREEN_FG, bold=True, center=True)  # stock receveur
            elif col == 10: sty(c, bg=bg, fg=RED_FG, bold=True, center=True)           # quantite
            elif col == 12: sty(c, bg=ORANGE_BG, fg=ORANGE_FG, bold=True, center=True) # seuil
            elif col == 13: sty(c, bg=GREEN_BG if actuel else ORANGE_BG,
                                fg=GREEN_FG if actuel else ORANGE_FG)
            else:           sty(c, bg=bg, fg=RED_FG)
        ws2.row_dimensions[row].height = 22
        row += 1

ws2.merge_cells(f'A{row}:I{row}')
ct = ws2.cell(row=row, column=1, value="TOTAL UNITES A TRANSFERER (CRITIQUE)")
ct.font = Font(bold=True, color=HEADER_FG, size=11, name='Arial')
ct.fill = PatternFill('solid', start_color=HEADER_BG)
ct.alignment = Alignment(horizontal='right', vertical='center')
ws2.row_dimensions[row].height = 25
hdr(ws2.cell(row=row, column=10, value=f"=SUM(J4:J{row-1})"))
ws2.freeze_panes = 'A4'

# ══════════════════════════════════════════════════════════════════════════
# FEUILLE 3 — ARTICLES FAIBLE (amelioree)
# ══════════════════════════════════════════════════════════════════════════
ws3 = wb.create_sheet("Faible - A surveiller")
ws3.sheet_view.showGridLines = False
nb_cols3 = 11

titre_principal(ws3, f"ARTICLES FAIBLE — A SURVEILLER   |   {today}", nb_cols3)
sous_titre(ws3, f"{len(faibles)} article(s) faible(s)   |   Action recommandee sous 1 semaine", nb_cols3)

hdrs3 = ['Code Article','Nom Article','Saison',
         'Score Donneur','De (Donneur)','Stock Donneur',
         'Score Receveur','Vers (Receveur)','Stock Receveur',
         'Quantite','Jours stock']
for col, h in enumerate(hdrs3, 1):
    hdr(ws3.cell(row=3, column=col, value=h), bg="B45309", fg="FFFFFF")
ws3.row_dimensions[3].height = 28

for i, w in enumerate([14,28,10,13,24,13,13,24,13,10,14], 1):
    ws3.column_dimensions[get_column_letter(i)].width = w

row = 4
for art in faibles:
    code   = str(art.get('codeArticle',''))
    nom    = noms_art.get(code,'')
    saison = art.get('saison','')
    stock_reel = stock_par_article.get(str(code), {})
    store_info = {b.get('storeName',''): b for b in art.get('analyse',[])}

    for s in art.get('suggestions', []):
        nd = s.get('de',''); nr = s.get('vers','')
        bd = store_info.get(nd, {}); br = store_info.get(nr, {})
        id_d = str(s.get('deId') or bd.get('storeId',''))
        id_r = str(s.get('versId') or br.get('storeId',''))
        sc_article = art.get('score', 0)
        sc_d = s.get('scoreDonneur', sc_article)
        sc_r = s.get('scoreReceveur', sc_article)
        jours_r = br.get('joursStock', 0)
        stk_d = max(0, stock_reel.get(id_d, 0) if isinstance(stock_reel.get(id_d, 0), (int,float)) else 0)
        stk_r = max(0, stock_reel.get(id_r, 0) if isinstance(stock_reel.get(id_r, 0), (int,float)) else 0)
        quantite = s.get('quantite', 0)
        if isinstance(quantite, (int, float)) and quantite < seuil:
            continue
        bg = ORANGE_BG if row % 2 == 0 else "FFE0B2"

        vals = [code, nom, saison, sc_d, nd, stk_d, sc_r, nr, stk_r,
                s.get('quantite',0), jours_r]

        for col, val in enumerate(vals, 1):
            c = ws3.cell(row=row, column=col, value=val)
            if   col == 1:  sty(c, bg=bg, fg=ORANGE_FG, bold=True)
            elif col == 2:  sty(c, bg=bg, fg="1E293B")
            elif col == 4:  sty(c, bg=BLUE_BG, fg=BLUE_FG, bold=True, center=True)
            elif col == 5:  sty(c, bg=BLUE_BG if nd=='SITE CENTRALE' else bg,
                                fg=BLUE_FG if nd=='SITE CENTRALE' else ORANGE_FG,
                                bold=(nd=='SITE CENTRALE'))
            elif col == 6:  sty(c, bg=BLUE_BG, fg=BLUE_FG, bold=True, center=True)
            elif col == 7:  sty(c, bg=GREEN_BG, fg=GREEN_FG, bold=True, center=True)
            elif col == 8:  sty(c, bg=GREEN_BG, fg=GREEN_FG)
            elif col == 9:  sty(c, bg=GREEN_BG, fg=GREEN_FG, bold=True, center=True)
            elif col == 10: sty(c, bg=bg, fg=ORANGE_FG, bold=True, center=True)
            else:           sty(c, bg=bg, fg=ORANGE_FG)
        ws3.row_dimensions[row].height = 22
        row += 1

ws3.freeze_panes = 'A4'

# ══════════════════════════════════════════════════════════════════════════
# FEUILLE 4 — VENTES PAR BOUTIQUE
# ══════════════════════════════════════════════════════════════════════════
ws4 = wb.create_sheet("Ventes par boutique")
ws4.sheet_view.showGridLines = False
nb_cols4 = 5
evol      = ventesHebdo.get('evolution', 0)
cette_sem = ventesHebdo.get('cetteSemaine', 0)
prec_sem  = ventesHebdo.get('semainePrec', 0)

titre_principal(ws4, f"VENTES HEBDOMADAIRES PAR BOUTIQUE   |   {today}", nb_cols4)
sous_titre(ws4, f"Cette semaine : {cette_sem} unites   |   Semaine precedente : {prec_sem} unites   |   Evolution : {evol:+}%", nb_cols4)

for col, h in enumerate(['Boutique (Store ID)','Unites vendues','References actives','Part du total (%)','Performance'], 1):
    hdr(ws4.cell(row=3, column=col, value=h))
ws4.row_dimensions[3].height = 22
for i, w in enumerate([24,18,18,16,16], 1):
    ws4.column_dimensions[get_column_letter(i)].width = w

parBoutique  = ventesHebdo.get('parBoutique', [])
total_ventes = sum(b.get('total_ventes',0) for b in parBoutique) or 1

for i, b in enumerate(parBoutique):
    row = 4 + i
    ws4.row_dimensions[row].height = 22
    bg     = GREY_BG if i % 2 == 0 else "FFFFFF"
    ventes = b.get('total_ventes', 0)
    part   = round((ventes / total_ventes) * 100, 1)
    if i < 3:                        perf_bg, perf_fg, perf_txt = GREEN_BG, GREEN_FG, "Top ventes"
    elif i >= len(parBoutique) - 3:  perf_bg, perf_fg, perf_txt = RED_BG,   RED_FG,   "Faibles ventes"
    else:                            perf_bg, perf_fg, perf_txt = bg,        "1E293B", "Moyen"
    sty(ws4.cell(row=row, column=1, value=b.get('store_id','')), bg=bg, bold=(i<3))
    sty(ws4.cell(row=row, column=2, value=ventes), bg=bg, bold=True, center=True)
    sty(ws4.cell(row=row, column=3, value=b.get('nb_references',0)), bg=bg, center=True)
    sty(ws4.cell(row=row, column=4, value=part), bg=bg, center=True)
    sty(ws4.cell(row=row, column=5, value=perf_txt), bg=perf_bg, fg=perf_fg, center=True, bold=(i<3))

row_tot = 4 + len(parBoutique)
ct4 = ws4.cell(row=row_tot, column=1, value="TOTAL GENERAL")
ct4.font = Font(bold=True, color=HEADER_FG, size=11, name='Arial')
ct4.fill = PatternFill('solid', start_color=HEADER_BG)
ct4.alignment = Alignment(horizontal='right', vertical='center')
ct4.border = brd()
ws4.row_dimensions[row_tot].height = 25
hdr(ws4.cell(row=row_tot, column=2, value=f"=SUM(B4:B{row_tot-1})"))
for col in range(3, 6):
    cd = ws4.cell(row=row_tot, column=col)
    cd.fill = PatternFill('solid', start_color=HEADER_BG); cd.border = brd()
ws4.freeze_panes = 'A4'

# ══════════════════════════════════════════════════════════════════════════
# FEUILLE 5 — ARTICLES INACTIFS (amelioree : + stock disponible + nom)
# ══════════════════════════════════════════════════════════════════════════
ws5 = wb.create_sheet("Articles inactifs")
ws5.sheet_view.showGridLines = False
nb_cols5 = 9
titre_principal(ws5, f"ARTICLES SANS MOUVEMENT CETTE SEMAINE   |   {today}", nb_cols5)
sous_titre(ws5, f"{len(inactifs)} article(s) actifs sur la periode mais sans vente cette semaine", nb_cols5)
hdrs5 = ['Code Article','Nom Article','Saison','Prix TTC','Ventes/sem moy.','Jours sans vente','Stock total','Urgence','Recommandation']
for col, h in enumerate(hdrs5, 1):
    hdr(ws5.cell(row=3, column=col, value=h), bg=PURPLE_FG, fg="FFFFFF")
ws5.row_dimensions[3].height = 28
for i, w in enumerate([14,28,10,10,14,16,12,12,42], 1):
    ws5.column_dimensions[get_column_letter(i)].width = w

def recommandation_inactif(saison, stock, jours_sans_vente, ventes_sem, prix):
    actuel = saison.strip().upper() in SAISONS_ACT
    if '26' in saison.upper():
        return ('NOUVEAU', 'Exposition insuffisante - laisser en rayon', '4472C4', 'D6E4F7')
    if actuel:
        if stock > 30 and jours_sans_vente > 7:
            return ('!! URGENT', 'Regrouper dans boutiques fortes - verifier visibilite', RED_FG, RED_BG)
        if stock > 10:
            return ('>> SURVEILLER', 'Verifier visibilite / repositionner en boutique', ORANGE_FG, ORANGE_BG)
        return ('OK', 'Stock faible - ecoulement naturel proche', GREEN_FG, GREEN_BG)
    else:
        if stock > 50:
            return ('!! DEMARQUE URGENTE', 'Stock eleve ancienne saison - demarque immediate', RED_FG, RED_BG)
        if stock > 20 and jours_sans_vente > 14:
            return ('>> DEMARQUE', 'Envisager demarque - article en fin de cycle', ORANGE_FG, ORANGE_BG)
        if stock > 0:
            return ('~ LIQUIDATION', 'Laisser ecouler - ne pas reapprovisionner', '7B6000', 'FFF9C4')
        return ('EPUISE', 'Stock nul - aucune action requise', '424242', 'F5F5F5')
for i, art in enumerate(inactifs):
    row = 4 + i
    ws5.row_dimensions[row].height = 24
    code          = str(art.get('code_article',''))
    nom           = noms_art.get(code,'')
    saison        = art.get('saison','')
    ventes_sem    = art.get('ventes_semaine', 0) or 0
    jours_sv      = art.get('jours_sans_vente', 0) or 0
    stock_reel    = stock_par_article.get(str(code), {})
    stock_total   = sum(v for v in stock_reel.values() if isinstance(v,(int,float)))
    art_db        = next((a for a in data.get('nomsArticles', {}).items() if a[0] == code), None)
    prix = prix_articles.get(str(code), 0)
    bg = PURPLE_BG if i % 2 == 0 else "EDE7F6"
    urgence_label, reco_text, reco_fg, reco_bg = recommandation_inactif(saison, stock_total, jours_sv, ventes_sem, prix)
    sty(ws5.cell(row=row, column=1, value=code), bg=bg, fg=PURPLE_FG, bold=True)
    sty(ws5.cell(row=row, column=2, value=nom), bg=bg, fg="1E293B")
    sty(ws5.cell(row=row, column=3, value=saison), bg=bg, fg=PURPLE_FG, center=True)
    sty(ws5.cell(row=row, column=4, value=prix if prix else 0), bg=bg, center=True)
    sty(ws5.cell(row=row, column=5, value=round(ventes_sem,2) if ventes_sem else '�'), bg=bg, fg=PURPLE_FG, center=True)
    # Jours sans vente avec couleur selon urgence
    c_jours = ws5.cell(row=row, column=6, value=jours_sv if jours_sv else '�')
    jours_bg = RED_BG if jours_sv > 14 else (ORANGE_BG if jours_sv > 7 else GREEN_BG)
    jours_fg = RED_FG if jours_sv > 14 else (ORANGE_FG if jours_sv > 7 else GREEN_FG)
    sty(c_jours, bg=jours_bg, fg=jours_fg, bold=True, center=True)
    # Stock avec couleur
    stock_bg = RED_BG if stock_total > 50 else (ORANGE_BG if stock_total > 20 else GREEN_BG)
    stock_fg = RED_FG if stock_total > 50 else (ORANGE_FG if stock_total > 20 else GREEN_FG)
    sty(ws5.cell(row=row, column=7, value=stock_total), bg=stock_bg, fg=stock_fg, bold=True, center=True)
    sty(ws5.cell(row=row, column=8, value=urgence_label), bg=reco_bg, fg=reco_fg, bold=True, center=True)
    sty(ws5.cell(row=row, column=9, value=reco_text), bg=reco_bg, fg=reco_fg)
ws5.freeze_panes = 'A4'

ws6 = wb.create_sheet("Score magasins")
ws6.sheet_view.showGridLines = False
nb_cols6 = 7

titre_principal(ws6, f"SCORE GLOBAL PAR MAGASIN   |   {today}", nb_cols6)
sous_titre(ws6, "Score = % articles critiques sur total articles exposes — Plus le score est bas, mieux c'est", nb_cols6)

for col, h in enumerate(['Store ID','Boutique','Articles exposes','Articles critiques','Articles faibles','Score (%)','Performance'], 1):
    hdr(ws6.cell(row=3, column=col, value=h))
ws6.row_dimensions[3].height = 22
for i, w in enumerate([12,28,18,18,16,12,18], 1):
    ws6.column_dimensions[get_column_letter(i)].width = w

if scores_mag:
    sorted_scores = sorted(scores_mag, key=lambda x: x.get('scorePct',0), reverse=True)
    for i, mg in enumerate(sorted_scores):
        row = 4 + i
        ws6.row_dimensions[row].height = 22
        pct  = mg.get('scorePct', 0)
        sid  = str(mg.get('storeId',''))
        if pct >= 30:   bg, fg, perf = RED_BG,    RED_FG,    "CRITIQUE"
        elif pct >= 15: bg, fg, perf = ORANGE_BG, ORANGE_FG, "ATTENTION"
        elif pct >= 5:  bg, fg, perf = GREY_BG,   "1E293B",  "OK"
        else:           bg, fg, perf = GREEN_BG,  GREEN_FG,  "BON"
        sty(ws6.cell(row=row, column=1, value=sid), bg=bg, center=True)
        sty(ws6.cell(row=row, column=2, value=STORE_NAMES.get(sid, mg.get('storeName',''))), bg=bg, fg=fg, bold=(pct>=30))
        sty(ws6.cell(row=row, column=3, value=mg.get('nbArticlesExposes',0)), bg=bg, center=True)
        sty(ws6.cell(row=row, column=4, value=mg.get('nbCritiques',0)),
            bg=RED_BG if mg.get('nbCritiques',0)>0 else bg,
            fg=RED_FG if mg.get('nbCritiques',0)>0 else "1E293B", bold=True, center=True)
        sty(ws6.cell(row=row, column=5, value=mg.get('nbFaibles',0)),
            bg=ORANGE_BG if mg.get('nbFaibles',0)>0 else bg,
            fg=ORANGE_FG if mg.get('nbFaibles',0)>0 else "1E293B", center=True)
        sty(ws6.cell(row=row, column=6, value=round(pct,1)), bg=bg, fg=fg, bold=True, center=True)
        sty(ws6.cell(row=row, column=7, value=perf), bg=bg, fg=fg, bold=(pct>=30), center=True)
else:
    ws6.merge_cells('A4:G4')
    c_nd = ws6.cell(row=4, column=1, value="Donnees non disponibles — ajouter 'scoresMagasins' dans le JSON retourne par l'API")
    c_nd.font = Font(color=ORANGE_FG, size=10, name='Arial', italic=True)
    c_nd.fill = PatternFill('solid', start_color=ORANGE_BG)
    c_nd.alignment = Alignment(horizontal='center', vertical='center')
    ws6.row_dimensions[4].height = 28

ws6.freeze_panes = 'A4'

# ══════════════════════════════════════════════════════════════════════════
# FEUILLE 7 — RUPTURES NOUVELLE COLLECTION 26E + 25H (NOUVELLE)
# ══════════════════════════════════════════════════════════════════════════
ws7 = wb.create_sheet("Ruptures nouvelle collection")
ws7.sheet_view.showGridLines = False
nb_cols7 = 8

titre_principal(ws7, f"RUPTURES DE STOCK — NOUVELLE COLLECTION (26E + 25H)   |   {today}", nb_cols7)
sous_titre(ws7, f"{len(rupturesNC)} article(s) en rupture ou quasi-rupture — Commande urgente recommandee", nb_cols7)

for col, h in enumerate(['Code Article','Nom Article','Saison','Boutique','Stock actuel','Ventes/sem','Jours restants','Action'], 1):
    hdr(ws7.cell(row=3, column=col, value=h), bg="B71C1C", fg="FFFFFF")
ws7.row_dimensions[3].height = 22
for i, w in enumerate([14,28,10,24,13,12,14,38], 1):
    ws7.column_dimensions[get_column_letter(i)].width = w

if rupturesNC:
    row = 4
    for art in rupturesNC:
        code   = str(art.get('code_article',''))
        nom    = noms_art.get(code, art.get('nom_article',''))
        saison = art.get('saison','')
        boutiques = art.get('boutiques', [])
        if not boutiques:
            boutiques = [{'storeName': art.get('storeName',''), 'stock': art.get('stock',0),
                          'ventesParSemaine': art.get('ventesParSemaine',0),
                          'joursRestants': art.get('joursRestants',0)}]
        for b in boutiques:
            stock = b.get('stock', 0)
            jours = b.get('joursRestants', 0)
            vps   = b.get('ventesParSemaine', 0)
            if stock == 0:    bg, fg, action = "FFCDD2","B71C1C","RUPTURE TOTALE — Commander immediatement"
            elif jours <= 7:  bg, fg, action = RED_BG,  RED_FG,  "Rupture imminente — Commander cette semaine"
            elif jours <= 14: bg, fg, action = ORANGE_BG,ORANGE_FG,"Stock faible — Preparer commande"
            else:             bg, fg, action = GREY_BG, "1E293B","A surveiller"
            ws7.row_dimensions[row].height = 22
            sty(ws7.cell(row=row, column=1, value=code), bg=bg, fg=fg, bold=True)
            sty(ws7.cell(row=row, column=2, value=nom), bg=bg, fg="1E293B")
            sty(ws7.cell(row=row, column=3, value=saison), bg=bg, fg=fg, center=True)
            sty(ws7.cell(row=row, column=4, value=b.get('storeName','')), bg=bg, fg=fg)
            sty(ws7.cell(row=row, column=5, value=stock),
                bg="FFCDD2" if stock==0 else bg, fg="B71C1C" if stock==0 else fg, bold=True, center=True)
            sty(ws7.cell(row=row, column=6, value=round(vps,1)), bg=bg, center=True)
            sty(ws7.cell(row=row, column=7, value=jours if jours>0 else '—'),
                bg=bg, fg=fg, bold=(jours<=7), center=True)
            sty(ws7.cell(row=row, column=8, value=action), bg=bg, fg=fg, bold=(stock==0))
            row += 1
else:
    ws7.merge_cells('A4:H4')
    c_ok = ws7.cell(row=4, column=1,
        value="Aucune rupture detectee — ajouter 'rupturesNouvelleCollection' dans l'API pour activer cette feuille")
    c_ok.font = Font(color=GREEN_FG, size=10, name='Arial', italic=True)
    c_ok.fill = PatternFill('solid', start_color=GREEN_BG)
    c_ok.alignment = Alignment(horizontal='center', vertical='center')
    ws7.row_dimensions[4].height = 30

ws7.freeze_panes = 'A4'

# ── SAUVEGARDER ────────────────────────────────────────────────────────────

# --------------------------------------------------------------------------
# FEUILLE 8 � BONS DE TRANSFERT (detail par variante taille/couleur)
# --------------------------------------------------------------------------
bons = data.get('bonsTransfert', [])
ws8 = wb.create_sheet("Bons de transfert")
ws8.sheet_view.showGridLines = False
nb_cols8 = 10
titre_principal(ws8, f"BONS DE TRANSFERT � DETAIL PAR VARIANTE   |   {today}", nb_cols8)
sous_titre(ws8, f"{len(bons)} transfert(s) a effectuer � Detail taille et couleur par article", nb_cols8)

hdrs8 = ['N� Bon','Code Article','Nom Article','Saison','EAN','Taille','Couleur','Stock Donneur','Stock Receveur','Qte a transferer']
for col, h in enumerate(hdrs8, 1):
    hdr(ws8.cell(row=3, column=col, value=h), bg=HEADER_BG, fg=HEADER_FG)
ws8.row_dimensions[3].height = 28
for i, w in enumerate([8,14,28,10,18,10,10,14,14,16], 1):
    ws8.column_dimensions[get_column_letter(i)].width = w

row = 4
bon_num = 1
for bon in bons:
    code     = str(bon.get('codeArticle',''))
    nom      = bon.get('nomArticle','')
    saison   = bon.get('saison','')
    donneur  = bon.get('donneur','')
    receveur = bon.get('receveur','')
    lignes   = bon.get('lignes', [])
    total    = bon.get('totalUnites', 0)

    # Entete du bon
    ws8.merge_cells(f'A{row}:J{row}')
    c_bon = ws8.cell(row=row, column=1,
        value=f"BON #{bon_num:03d}  |  DE : {donneur}  ->  VERS : {receveur}  |  {code} {nom}  |  TOTAL : {total} unites")
    c_bon.font = Font(bold=True, color=HEADER_FG, size=11, name='Arial')
    c_bon.fill = PatternFill('solid', start_color=HEADER_BG)
    c_bon.alignment = Alignment(horizontal='left', vertical='center', indent=1)
    ws8.row_dimensions[row].height = 28
    row += 1

    # Lignes variantes
    for j, ligne in enumerate(lignes):
        bg = "EBF5FB" if j % 2 == 0 else "FFFFFF"
        qte = ligne.get('quantite', 0)
        stk_d = ligne.get('stockDonneur', 0)
        stk_r = ligne.get('stockReceveur', 0)
        vals = [
            f"#{bon_num:03d}", code, nom, saison,
            ligne.get('ean',''), ligne.get('taille',''), ligne.get('couleur',''),
            stk_d, stk_r, qte
        ]
        for col, val in enumerate(vals, 1):
            c = ws8.cell(row=row, column=col, value=val)
            if col == 10:
                sty(c, bg=GREEN_BG, fg=GREEN_FG, bold=True, center=True)
            elif col in [8, 9]:
                stk_bg = RED_BG if (col==9 and stk_r==0) else bg
                stk_fg = RED_FG if (col==9 and stk_r==0) else "1E293B"
                sty(c, bg=stk_bg, fg=stk_fg, bold=(col==9 and stk_r==0), center=True)
            elif col in [5, 6, 7]:
                sty(c, bg=bg, fg="1E293B", center=True)
            else:
                sty(c, bg=bg, fg="1E293B")
        ws8.row_dimensions[row].height = 22
        row += 1

    # Ligne total par bon
    ws8.merge_cells(f'A{row}:I{row}')
    ct = ws8.cell(row=row, column=1, value=f"TOTAL BON #{bon_num:03d} -- {donneur} -> {receveur}")
    ct.font = Font(bold=True, color=HEADER_FG, size=10, name='Arial')
    ct.fill = PatternFill('solid', start_color=HEADER_BG)
    ct.alignment = Alignment(horizontal='right', vertical='center')
    ws8.row_dimensions[row].height = 22
    hdr(ws8.cell(row=row, column=10, value=total))
    row += 2
    bon_num += 1

if not bons:
    ws8.merge_cells('A4:J4')
    c_ok = ws8.cell(row=4, column=1, value="Aucun transfert detaille disponible")
    c_ok.font = Font(color=GREEN_FG, size=10, name='Arial', italic=True)
    c_ok.fill = PatternFill('solid', start_color=GREEN_BG)
    c_ok.alignment = Alignment(horizontal='center', vertical='center')
    ws8.row_dimensions[4].height = 30
ws8.freeze_panes = 'A4'



# --------------------------------------------------------------------------
# FEUILLE 9 � RETOURS & ANNULATIONS
# --------------------------------------------------------------------------
analyse_retours = data.get('analyseRetours', {})
retour_stats    = analyse_retours.get('parArticle', [])
retour_boutique = analyse_retours.get('parBoutique', [])
taux_retour     = analyse_retours.get('tauxRetour', [])
retour_total    = analyse_retours.get('total', {})

ws9 = wb.create_sheet("Retours & Annulations")
ws9.sheet_view.showGridLines = False
nb_cols9 = 8
titre_principal(ws9, f"RETOURS & ANNULATIONS   |   {today}", nb_cols9)
total_retours = retour_total.get('total', 0) if retour_total else 0
sous_titre(ws9, f"{total_retours} unites retournees sur la periode � Analyse impact et articles a risque", nb_cols9)

row = 3

# === SECTION 1 : TOP ARTICLES RETOURNES ===
ws9.merge_cells(f'A{row}:H{row}')
c = ws9.cell(row=row, column=1, value="TOP ARTICLES RETOURNES � TAUX DE RETOUR PAR ARTICLE")
c.font = Font(bold=True, color=RED_FG, size=11, name='Arial')
c.fill = PatternFill('solid', start_color=RED_BG)
c.alignment = Alignment(horizontal='center', vertical='center')
ws9.row_dimensions[row].height = 28
row += 1

hdrs_r1 = ['Code Article','Nom Article','Famille','Prix TTC','Ventes 28j','Retours','Taux retour (%)','Alerte']
for col, h in enumerate(hdrs_r1, 1):
    hdr(ws9.cell(row=row, column=col, value=h), bg=RED_FG, fg="FFFFFF")
ws9.row_dimensions[row].height = 22
row += 1

for i, w in enumerate([14,28,16,10,12,10,14,20], 1):
    ws9.column_dimensions[get_column_letter(i)].width = w

for art in taux_retour:
    code    = str(art.get('code_article',''))
    nom     = art.get('libelle','') or noms_art.get(code,'')
    famille = art.get('famille','')
    prix    = art.get('prix_detail', 0) or prix_articles.get(code, 0)
    ventes  = art.get('ventes', 0)
    retours = art.get('retours', 0)
    taux    = art.get('taux_pct', 0)
    
    if taux >= 20:
        alerte = "!! PROBLEME QUALITE/TAILLE"
        bg_row = RED_BG; fg_row = RED_FG
    elif taux >= 10:
        alerte = ">> A SURVEILLER"
        bg_row = ORANGE_BG; fg_row = ORANGE_FG
    else:
        alerte = "OK"
        bg_row = "F1F8E9"; fg_row = GREEN_FG

    vals = [code, nom, famille, prix, ventes, retours, taux, alerte]
    for col, val in enumerate(vals, 1):
        c = ws9.cell(row=row, column=col, value=val)
        if col == 7:
            taux_bg = RED_BG if taux >= 20 else (ORANGE_BG if taux >= 10 else "F1F8E9")
            taux_fg = RED_FG if taux >= 20 else (ORANGE_FG if taux >= 10 else GREEN_FG)
            sty(c, bg=taux_bg, fg=taux_fg, bold=True, center=True)
        elif col == 8:
            sty(c, bg=bg_row, fg=fg_row, bold=(taux>=10))
        else:
            sty(c, bg=bg_row, fg=fg_row, center=(col in [4,5,6]))
    ws9.row_dimensions[row].height = 22
    row += 1

row += 1

# === SECTION 2 : RETOURS PAR BOUTIQUE ===
ws9.merge_cells(f'A{row}:H{row}')
c2 = ws9.cell(row=row, column=1, value="RETOURS PAR BOUTIQUE")
c2.font = Font(bold=True, color=ORANGE_FG, size=11, name='Arial')
c2.fill = PatternFill('solid', start_color=ORANGE_BG)
c2.alignment = Alignment(horizontal='center', vertical='center')
ws9.row_dimensions[row].height = 28
row += 1

STORE_NAMES = {
    '002':'MENZAH','005':'JAMEL ABDENNACEUR','009':'SFAX A','011':'NABEUL',
    '014':'CARREFOUR','015':'ZEPHYR','016':'LAFAYETTE','019':'E-Boutique',
    '021':'TUNISIA MALL','024':'GEANT 3','029':'SOUSSE SLIM CENTER',
    '030':'AZUR CITY','031':'SOUKRA','032':'MALL OF SFAX','033':'LAC Premium'
}

hdrs_r2 = ['Store ID','Boutique','Unites retournees','Articles concernes','','','','']
for col, h in enumerate(hdrs_r2[:4], 1):
    hdr(ws9.cell(row=row, column=col, value=h), bg=ORANGE_FG, fg="FFFFFF")
ws9.row_dimensions[row].height = 22
row += 1

max_retours = max((b.get('total',0) for b in retour_boutique), default=1)
for b in retour_boutique:
    store_id = str(b.get('store_id',''))
    nom_store = STORE_NAMES.get(store_id, store_id)
    total_b   = b.get('total', 0)
    nb_art    = b.get('nb_articles', b.get('nb', 0))
    pct       = round(total_b * 100 / max_retours) if max_retours > 0 else 0
    bg_b = RED_BG if pct >= 80 else (ORANGE_BG if pct >= 50 else "FFF9C4")
    fg_b = RED_FG if pct >= 80 else (ORANGE_FG if pct >= 50 else "7B6000")
    sty(ws9.cell(row=row, column=1, value=store_id), bg=bg_b, fg=fg_b, center=True)
    sty(ws9.cell(row=row, column=2, value=nom_store), bg=bg_b, fg=fg_b, bold=True)
    sty(ws9.cell(row=row, column=3, value=total_b), bg=bg_b, fg=fg_b, bold=True, center=True)
    sty(ws9.cell(row=row, column=4, value=nb_art), bg=bg_b, fg=fg_b, center=True)
    ws9.row_dimensions[row].height = 22
    row += 1

ws9.freeze_panes = 'A4'


wb.save(output_path)
print(f"Rapport genere : {output_path}")